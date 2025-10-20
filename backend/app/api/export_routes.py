"""Export API Routes - Exportação de dados em múltiplos formatos abrangentes."""

from datetime import datetime, timedelta
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
import csv
import io
import json

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.orm import Session, joinedload, selectinload

from app.db.session import SessionLocal
from app.models.models import Citation, Entity, Run, SerpFeature, RunSemanticInsight
from app.services.semantic_payload import load_insight_payload

router = APIRouter(tags=["export"])


def get_db() -> Iterable[Session]:
    """Database session dependency."""
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _parse_engine_ids(engine_ids: Optional[str]) -> List[str]:
    if not engine_ids:
        return []
    return [value.strip() for value in engine_ids.split(",") if value.strip()]


def _apply_run_filters(
    db: Session,
    query,
    *,
    project_id: Optional[str],
    engine_ids: Optional[str],
    subproject_id: Optional[str],
    prompt_id: Optional[str],
    run_id: Optional[str],
    im_seo_min: Optional[float],
    im_seo_max: Optional[float],
    im_seoia_min: Optional[float],
    im_seoia_max: Optional[float],
):
    engine_id_list = _parse_engine_ids(engine_ids)

    if project_id:
        query = query.filter(Run.project_id == project_id)

    if engine_id_list:
        query = query.filter(Run.engine_id.in_(engine_id_list))

    if subproject_id:
        query = query.filter(Run.subproject_id == subproject_id)

    if prompt_id:
        from app.models.models import PromptVersion

        prompt_version_ids = (
            db.query(PromptVersion.id)
            .filter(PromptVersion.prompt_id == prompt_id)
            .subquery()
        )
        query = query.filter(Run.prompt_version_id.in_(prompt_version_ids))

    if run_id:
        query = query.filter(Run.id == run_id)

    if im_seo_min is not None:
        query = query.filter(Run.im_seo_score >= im_seo_min)
    if im_seo_max is not None:
        query = query.filter(Run.im_seo_score <= im_seo_max)

    if im_seoia_min is not None:
        query = query.filter(Run.im_seoia_score >= im_seoia_min)
    if im_seoia_max is not None:
        query = query.filter(Run.im_seoia_score <= im_seoia_max)

    return query


def _resolve_period(days: int, start_date: Optional[str], end_date: Optional[str]) -> Tuple[datetime, datetime]:
    if start_date and end_date:
        try:
            return datetime.fromisoformat(start_date), datetime.fromisoformat(end_date)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="Datas inválidas. Use ISO-8601.") from exc
    end = datetime.utcnow()
    start = end - timedelta(days=days)
    return start, end


def _collect_runs(
    db: Session,
    *,
    project_id: Optional[str],
    engine_ids: Optional[str],
    days: int,
    start_date: Optional[str],
    end_date: Optional[str],
    subproject_id: Optional[str],
    prompt_id: Optional[str],
    run_id: Optional[str],
    im_seo_min: Optional[float],
    im_seo_max: Optional[float],
    im_seoia_min: Optional[float],
    im_seoia_max: Optional[float],
) -> Tuple[List[Run], Tuple[datetime, datetime]]:
    cutoff_start, cutoff_end = _resolve_period(days, start_date, end_date)

    query = (
        db.query(Run)
        .options(
            selectinload(Run.citations),
            selectinload(Run.entities),
            selectinload(Run.serp_features),
            selectinload(Run.semantic_insights),
        )
        .filter(
            Run.status == "completed",
            Run.started_at >= cutoff_start,
            Run.started_at <= cutoff_end,
        )
    )

    query = _apply_run_filters(
        db,
        query,
        project_id=project_id,
        engine_ids=engine_ids,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )

    runs: List[Run] = query.order_by(Run.started_at.desc()).all()
    return runs, (cutoff_start, cutoff_end)


def _hydrate_references(db: Session, runs: Sequence[Run]) -> Dict[str, Dict[str, Optional[str]]]:
    from app.models.models import Engine, Project, Prompt, PromptVersion, SubProject

    project_ids = {run.project_id for run in runs if run.project_id}
    subproject_ids = {run.subproject_id for run in runs if run.subproject_id}
    engine_ids = {run.engine_id for run in runs if run.engine_id}
    prompt_version_ids = {run.prompt_version_id for run in runs if run.prompt_version_id}

    prompt_ids = set()
    if prompt_version_ids:
        prompt_versions = {
            pv.id: pv
            for pv in db.query(PromptVersion).filter(PromptVersion.id.in_(prompt_version_ids)).all()
        }
        prompt_ids = {pv.prompt_id for pv in prompt_versions.values() if pv.prompt_id}
    else:
        prompt_versions = {}

    prompts = {
        prompt.id: prompt
        for prompt in db.query(Prompt).filter(Prompt.id.in_(prompt_ids)).all()
    } if prompt_ids else {}

    engines = {
        engine.id: engine
        for engine in db.query(Engine).filter(Engine.id.in_(engine_ids)).all()
    } if engine_ids else {}

    projects = {
        project.id: project
        for project in db.query(Project).filter(Project.id.in_(project_ids)).all()
    } if project_ids else {}

    subprojects = {
        subproject.id: subproject
        for subproject in db.query(SubProject).filter(SubProject.id.in_(subproject_ids)).all()
    } if subproject_ids else {}

    return {
        "prompt_versions": prompt_versions,
        "prompts": prompts,
        "engines": engines,
        "projects": projects,
        "subprojects": subprojects,
    }


def _run_base_row(run: Run, refs: Dict[str, Dict[str, Optional[str]]]) -> Dict[str, Optional[object]]:
    prompt_version = refs["prompt_versions"].get(run.prompt_version_id)
    prompt = refs["prompts"].get(prompt_version.prompt_id) if prompt_version else None
    engine = refs["engines"].get(run.engine_id)
    project = refs["projects"].get(run.project_id)
    subproject = refs["subprojects"].get(run.subproject_id)

    citations_is_ours = any(citation.is_ours for citation in getattr(run, "citations", []) or [])

    base_row: Dict[str, Optional[object]] = {
        "id": run.id,
        "project_id": run.project_id,
        "project_name": getattr(project, "name", None),
        "subproject_id": run.subproject_id,
        "subproject_name": getattr(subproject, "name", None),
        "prompt_version_id": run.prompt_version_id,
        "prompt_id": getattr(prompt_version, "prompt_id", None),
        "prompt_name": getattr(prompt, "name", None),
        "prompt_version": getattr(prompt_version, "version", None),
        "engine_id": run.engine_id,
        "engine_name": getattr(engine, "name", None),
        "engine_region": getattr(engine, "region", None),
        "engine_device": getattr(engine, "device", None),
        "status": run.status,
        "monitor_id": run.monitor_id,
        "model_name": run.model_name,
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        "latency_ms": run.latency_ms,
        "cost_usd": run.cost_usd,
        "tokens_input": run.tokens_input,
        "tokens_output": run.tokens_output,
        "tokens_total": run.tokens_total,
        "cycles_total": run.cycles_total,
        "cycle_delay_seconds": run.cycle_delay_seconds,
        "citations_count": run.citations_count,
        "our_citations_count": run.our_citations_count,
        "citations_is_ours": citations_is_ours,
        "unique_domains_count": run.unique_domains_count,
        "im_seo_score": run.im_seo_score,
        "im_seoia_score": run.im_seoia_score,
        "eeat_score": run.eeat_score,
        "eeat_expertise": run.eeat_expertise,
        "eeat_experience": run.eeat_experience,
        "eeat_authoritativeness": run.eeat_authoritativeness,
        "eeat_trustworthiness": run.eeat_trustworthiness,
        "core_web_vitals_score": run.core_web_vitals_score,
        "lcp_score": run.lcp_score,
        "fid_score": run.fid_score,
        "cls_score": run.cls_score,
        "irzc_score": run.irzc_score,
        "ctr_expected": run.ctr_expected,
        "ctr_real": run.ctr_real,
        "ctr_ratio": run.ctr_ratio,
        "share_of_voice_serp": run.share_of_voice_serp,
        "serp_features_presence": run.serp_features_presence,
        "ia_ready_score": run.ia_ready_score,
        "ia_ready_blocks_count": run.ia_ready_blocks_count,
        "has_lists": run.has_lists,
        "has_faqs": run.has_faqs,
        "has_tables": run.has_tables,
        "has_step_by_step": run.has_step_by_step,
        "entities_detected": run.entities_detected,
        "entities_relevance_score": run.entities_relevance_score,
        "entity_connection_score": run.entity_connection_score,
        "schema_types_detected": run.schema_types_detected,
        "schema_coverage_score": run.schema_coverage_score,
        "schema_valid": run.schema_valid,
        "response_type": run.response_type,
        "sufficiency_level": run.sufficiency_level,
        "actionability_type": run.actionability_type,
        "trust_source": run.trust_source,
        "brand_positioning": run.brand_positioning,
        "question_type": run.question_type,
        "funnel_stage": run.funnel_stage,
        "classification_confidence": run.classification_confidence,
        "classification_version": run.classification_version,
        "perceived_value_category": run.perceived_value_category,
        "semantic_summary": run.semantic_summary,
        "schedule_date": run.schedule_date.isoformat() if run.schedule_date else None,
        "schedule_slot": run.schedule_slot,
        "schedule_index_today": run.schedule_index_today,
        "schedule_total_today": run.schedule_total_today,
        "schedule_source": run.schedule_source,
        "response_text": run.response_text,
        "semantic_insights_updated_at": (
            run.semantic_insights.updated_at.isoformat() if getattr(run, "semantic_insights", None) else None
        ),
    }

    return base_row


def _run_citation_rows(run: Run) -> List[Dict[str, Optional[object]]]:
    rows: List[Dict[str, Optional[object]]] = []
    for citation in getattr(run, "citations", []) or []:
        rows.append(
            {
                "run_id": run.id,
                "citation_id": citation.id,
                "domain": citation.domain,
                "url": citation.url,
                "anchor": citation.anchor,
                "position": citation.position,
                "type": citation.type,
                "is_ours": citation.is_ours,
            }
        )
    return rows


def _run_serp_feature_rows(run: Run) -> List[Dict[str, Optional[object]]]:
    rows: List[Dict[str, Optional[object]]] = []
    for feature in getattr(run, "serp_features", []) or []:
        rows.append(
            {
                "run_id": run.id,
                "serp_feature_id": feature.id,
                "has_featured_snippet": feature.has_featured_snippet,
                "has_paa": feature.has_paa,
                "has_knowledge_panel": feature.has_knowledge_panel,
                "has_ai_overview": feature.has_ai_overview,
                "has_local_pack": feature.has_local_pack,
                "has_video_carousel": feature.has_video_carousel,
                "has_image_pack": feature.has_image_pack,
                "featured_snippet_content": feature.featured_snippet_content,
                "paa_questions": feature.paa_questions,
                "paa_items": feature.paa_items,
                "knowledge_panel_json": feature.knowledge_panel_json,
                "ai_overview_json": feature.ai_overview_json,
                "organic_position": feature.organic_position,
                "competitors_in_top10": feature.competitors_in_top10,
            }
        )
    return rows


def _run_entity_rows(run: Run) -> List[Dict[str, Optional[object]]]:
    rows: List[Dict[str, Optional[object]]] = []
    for entity in getattr(run, "entities", []) or []:
        rows.append(
            {
                "run_id": run.id,
                "entity_id": entity.id,
                "name": entity.name,
                "entity_type": entity.entity_type,
                "salience_score": entity.salience_score,
                "mentions_count": entity.mentions_count,
            }
        )
    return rows


def _build_metadata(
    *,
    runs: Sequence[Run],
    period: Tuple[datetime, datetime],
    filters: Dict[str, Optional[object]],
) -> Dict[str, object]:
    return {
        "exported_at": datetime.utcnow().isoformat(),
        "period": {
            "start": period[0].isoformat(),
            "end": period[1].isoformat(),
        },
        "total_runs": len(runs),
        "filters": filters,
        "metrics_summary": {
            "im_seo_avg": _safe_avg([run.im_seo_score for run in runs]),
            "im_seoia_avg": _safe_avg([run.im_seoia_score for run in runs]),
            "eeat_avg": _safe_avg([run.eeat_score for run in runs]),
            "core_web_vitals_avg": _safe_avg([run.core_web_vitals_score for run in runs]),
            "irzc_avg": _safe_avg([run.irzc_score for run in runs]),
        },
    }


def _safe_avg(values: Sequence[Optional[float]]) -> Optional[float]:
    filtered = [value for value in values if value is not None]
    if not filtered:
        return None
    return round(sum(filtered) / len(filtered), 2)


def _build_filters_dict(**kwargs: Optional[object]) -> Dict[str, Optional[object]]:
    return {key: value for key, value in kwargs.items() if value is not None and value != []}


@router.get("/export/csv")
def export_csv(
    project_id: Optional[str] = None,
    engine_ids: Optional[str] = None,
    days: int = Query(30, ge=1, le=365),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    include_metadata: bool = True,
    subproject_id: Optional[str] = None,
    prompt_id: Optional[str] = None,
    run_id: Optional[str] = None,
    im_seo_min: Optional[float] = None,
    im_seo_max: Optional[float] = None,
    im_seoia_min: Optional[float] = None,
    im_seoia_max: Optional[float] = None,
    db: Session = Depends(get_db),
):
    runs, period = _collect_runs(
        db,
        project_id=project_id,
        engine_ids=engine_ids,
        days=days,
        start_date=start_date,
        end_date=end_date,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )

    refs = _hydrate_references(db, runs)

    output = io.StringIO()

    if include_metadata:
        metadata = _build_metadata(
            runs=runs,
            period=period,
            filters=_build_filters_dict(
                project_id=project_id,
                engine_ids=engine_ids,
                subproject_id=subproject_id,
                prompt_id=prompt_id,
                run_id=run_id,
                days=days,
                start_date=start_date,
                end_date=end_date,
                im_seo_min=im_seo_min,
                im_seo_max=im_seo_max,
                im_seoia_min=im_seoia_min,
                im_seoia_max=im_seoia_max,
            ),
        )
        output.write(f"# metadata: {json.dumps(metadata, ensure_ascii=False)}\n")

    fieldnames = sorted(
        {
            key
            for run in runs
            for key in _run_base_row(run, refs).keys()
        }
    )
    fieldnames.extend(["citations_json", "serp_features_json", "entities_json", "semantic_insights_json"])

    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()

    for run in runs:
        base_row = _run_base_row(run, refs)
        base_row["citations_json"] = json.dumps(_run_citation_rows(run), ensure_ascii=False)
        base_row["serp_features_json"] = json.dumps(_run_serp_feature_rows(run), ensure_ascii=False)
        base_row["entities_json"] = json.dumps(_run_entity_rows(run), ensure_ascii=False)
        base_row["semantic_insights_json"] = json.dumps(
            getattr(run.semantic_insights, "payload", None), ensure_ascii=False
        ) if getattr(run, "semantic_insights", None) else None
        writer.writerow(base_row)

    filename = f"im_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"

    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
        },
    )


@router.get("/export/json")
def export_json(
    project_id: Optional[str] = None,
    engine_ids: Optional[str] = None,
    days: int = Query(30, ge=1, le=365),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    include_metadata: bool = True,
    subproject_id: Optional[str] = None,
    prompt_id: Optional[str] = None,
    run_id: Optional[str] = None,
    im_seo_min: Optional[float] = None,
    im_seo_max: Optional[float] = None,
    im_seoia_min: Optional[float] = None,
    im_seoia_max: Optional[float] = None,
    db: Session = Depends(get_db),
):
    runs, period = _collect_runs(
        db,
        project_id=project_id,
        engine_ids=engine_ids,
        days=days,
        start_date=start_date,
        end_date=end_date,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )

    refs = _hydrate_references(db, runs)

    runs_payload = []
    for run in runs:
        base_row = _run_base_row(run, refs)
        runs_payload.append(
            {
                **base_row,
                "citations": _run_citation_rows(run),
                "serp_features": _run_serp_feature_rows(run),
                "entities": _run_entity_rows(run),
            }
        )

    response_data: Dict[str, object] = {"runs": runs_payload}

    if include_metadata:
        response_data["metadata"] = _build_metadata(
            runs=runs,
            period=period,
            filters=_build_filters_dict(
                project_id=project_id,
                engine_ids=engine_ids,
                subproject_id=subproject_id,
                prompt_id=prompt_id,
                run_id=run_id,
                days=days,
                start_date=start_date,
                end_date=end_date,
                im_seo_min=im_seo_min,
                im_seo_max=im_seo_max,
                im_seoia_min=im_seoia_min,
                im_seoia_max=im_seoia_max,
            ),
        )

    filename = f"im_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.json"

    return Response(
        content=json.dumps(response_data, ensure_ascii=False, indent=2),
        media_type="application/json",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
        },
    )


@router.get("/export/excel")
def export_excel(
    project_id: Optional[str] = None,
    engine_ids: Optional[str] = None,
    days: int = Query(30, ge=1, le=365),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    subproject_id: Optional[str] = None,
    prompt_id: Optional[str] = None,
    run_id: Optional[str] = None,
    im_seo_min: Optional[float] = None,
    im_seo_max: Optional[float] = None,
    im_seoia_min: Optional[float] = None,
    im_seoia_max: Optional[float] = None,
    db: Session = Depends(get_db),
):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="Excel export requires 'openpyxl'. Execute: pip install openpyxl",
        ) from exc

    runs, period = _collect_runs(
        db,
        project_id=project_id,
        engine_ids=engine_ids,
        days=days,
        start_date=start_date,
        end_date=end_date,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )

    refs = _hydrate_references(db, runs)

    wb = openpyxl.Workbook()

    # Metadata sheet
    meta_ws = wb.active
    meta_ws.title = "Metadata"
    metadata = _build_metadata(
        runs=runs,
        period=period,
        filters=_build_filters_dict(
            project_id=project_id,
            engine_ids=engine_ids,
            subproject_id=subproject_id,
            prompt_id=prompt_id,
            run_id=run_id,
            days=days,
            start_date=start_date,
            end_date=end_date,
            im_seo_min=im_seo_min,
            im_seo_max=im_seo_max,
            im_seoia_min=im_seoia_min,
            im_seoia_max=im_seoia_max,
        ),
    )

    meta_ws.cell(row=1, column=1, value="Chave").font = Font(bold=True)
    meta_ws.cell(row=1, column=2, value="Valor").font = Font(bold=True)

    row_idx = 2
    for key, value in metadata.items():
        meta_ws.cell(row=row_idx, column=1, value=key)
        meta_ws.cell(row=row_idx, column=2, value=json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value)
        row_idx += 1

    for column_cells in meta_ws.columns:
        max_length = max(len(str(cell.value)) for cell in column_cells if cell.value)
        meta_ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 80)

    # Runs sheet
    runs_ws = wb.create_sheet("Runs")
    base_rows = [_run_base_row(run, refs) for run in runs]
    if base_rows:
        headers = list(base_rows[0].keys())
    else:
        headers = []

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, 1):
        cell = runs_ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row in enumerate(base_rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            runs_ws.cell(row=row_idx, column=col_idx, value=row.get(header))

    for column_cells in runs_ws.columns:
        max_length = max(len(str(cell.value)) for cell in column_cells if cell.value)
        runs_ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 80)

    # Citations sheet
    citations_ws = wb.create_sheet("Citations")
    citations_rows = [row for run in runs for row in _run_citation_rows(run)]
    _populate_sheet(citations_ws, citations_rows)

    # SERP Features sheet
    serp_ws = wb.create_sheet("SERP Features")
    serp_rows = [row for run in runs for row in _run_serp_feature_rows(run)]
    _populate_sheet(serp_ws, serp_rows)

    # Entities sheet
    entities_ws = wb.create_sheet("Entities")
    entity_rows = [row for run in runs for row in _run_entity_rows(run)]
    _populate_sheet(entities_ws, entity_rows)

    # Autosize alignment for detail sheets
    for ws in [citations_ws, serp_ws, entities_ws]:
        for column_cells in ws.columns:
            if not column_cells:
                continue
            max_length = max(len(str(cell.value)) for cell in column_cells if cell.value)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 80)
            for cell in column_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"im_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
        },
    )


def _populate_sheet(ws, rows: List[Dict[str, Optional[object]]]) -> None:
    if not rows:
        ws.append(["Nenhum registro disponível"])
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])


# ============================================================================
# HELPER FUNCTIONS FOR NEW GEO EXPORT TABLES
# ============================================================================

def _check_bb_mention(text: str) -> bool:
    """Verifica se texto menciona Banco do Brasil."""
    if not text:
        return False
    text_lower = text.lower()
    return any(term in text_lower for term in ["banco do brasil", " bb ", "banco brasil"])


def _find_bb_position(text: str) -> Optional[int]:
    """Encontra posição da primeira menção ao BB no texto."""
    if not text:
        return None
    text_lower = text.lower()
    for term in ["banco do brasil", "banco brasil", " bb "]:
        pos = text_lower.find(term)
        if pos != -1:
            return pos
    return None


def _parse_ai_overview(serp_feature) -> Dict[str, Any]:
    """Extrai e estrutura dados do AI Overview JSON."""
    if not serp_feature or not serp_feature.ai_overview_json:
        return {
            "summary": "",
            "text_blocks": [],
            "references": [],
            "follow_up_questions": []
        }
    
    try:
        data = json.loads(serp_feature.ai_overview_json)
        return {
            "summary": data.get("summary", ""),
            "text_blocks": data.get("text_blocks", []),
            "references": data.get("references", []),
            "follow_up_questions": data.get("follow_up_questions", [])
        }
    except (json.JSONDecodeError, TypeError):
        return {
            "summary": "",
            "text_blocks": [],
            "references": [],
            "follow_up_questions": []
        }


def _build_response_text(ai_overview: Dict[str, Any]) -> str:
    """Constrói texto completo da resposta a partir do AI Overview."""
    parts = [ai_overview["summary"]]
    
    for block in ai_overview["text_blocks"]:
        if block.get("type") == "paragraph":
            parts.append(block.get("snippet", ""))
        elif block.get("type") == "list":
            for item in block.get("list", []):
                parts.append(f"• {item.get('snippet', '')}")
    
    return "\n\n".join(filter(None, parts))


def _get_bb_url_info(references: List[Dict], citations: List) -> Dict[str, Any]:
    """Analisa URLs do BB nas referências e citações."""
    bb_refs = []
    for idx, ref in enumerate(references):
        domain = ref.get("domain", "").lower()
        is_bb = "bb.com.br" in domain or "banco do brasil" in domain.lower()
        bb_refs.append({
            "position": idx + 1,
            "domain": ref.get("domain"),
            "url": ref.get("url"),
            "is_bb": is_bb
        })
    
    # Verificar também citations
    for cite in citations:
        if cite.is_ours and not any(r["url"] == cite.url for r in bb_refs):
            bb_refs.append({
                "position": len(bb_refs) + 1,
                "domain": cite.domain,
                "url": cite.url,
                "is_bb": True
            })
    
    has_bb_url = any(r["is_bb"] for r in bb_refs)
    first_bb_pos = next((r["position"] for r in bb_refs if r["is_bb"]), None)
    citation_binary = "".join("1" if r["is_bb"] else "0" for r in bb_refs[:10])
    
    return {
        "tem_url": "Sim" if bb_refs else "Não",
        "url_bb": "Sim" if has_bb_url else "Não",
        "posicao_url_bb": first_bb_pos,
        "citation_bb": citation_binary if citation_binary else None,
        "ranking_url": json.dumps(bb_refs, ensure_ascii=False) if bb_refs else "[]"
    }


def _classify_relevance(im_seo_score: Optional[float], funnel_stage: Optional[str]) -> str:
    """Classifica nível de relevância."""
    if im_seo_score is None:
        return "Não classificado"
    
    if im_seo_score >= 80:
        return "Alta"
    elif im_seo_score >= 60:
        return "Média-Alta"
    elif im_seo_score >= 40:
        return "Média"
    else:
        return "Consultiva/Reputacional"


def _map_funnel_stage(stage: Optional[str]) -> str:
    """Mapeia etapa do funil."""
    if not stage:
        return "Não classificado"
    
    mapping = {
        "reconhecimento": "Reconhecimento",
        "consideracao": "Consideração",
        "conversao": "Conversão"
    }
    return mapping.get(stage.lower(), stage.capitalize())


def _calculate_perception_scores(perception: Dict[str, Any]) -> Dict[str, int]:
    """Calcula notas de percepção (0-5) baseado em categorias."""
    scores = {
        "inovador": 0,
        "seguranca": 0,
        "custo": 0,
        "atendimento": 0
    }
    
    primary = perception.get("primary_category", "").lower()
    secondary = [s.lower() for s in perception.get("secondary_categories", [])]
    
    # Categoria primária: 5 pontos
    if "inovacao" in primary or "inovador" in primary:
        scores["inovador"] = 5
    elif "tradicao" in primary or "seguranca" in primary:
        scores["seguranca"] = 5
    elif "custo" in primary or "baixo custo" in primary:
        scores["custo"] = 5
    elif "atendimento" in primary or "relacionamento" in primary:
        scores["atendimento"] = 5
    
    # Categorias secundárias: 3 pontos
    for cat in secondary:
        if "inovacao" in cat and scores["inovador"] < 3:
            scores["inovador"] = 3
        elif "tradicao" in cat or "seguranca" in cat and scores["seguranca"] < 3:
            scores["seguranca"] = 3
        elif "custo" in cat and scores["custo"] < 3:
            scores["custo"] = 3
        elif "atendimento" in cat and scores["atendimento"] < 3:
            scores["atendimento"] = 3
    
    return scores


# ============================================================================
# EXPORT ENDPOINT: TABELA SERP - AI OVERVIEW
# ============================================================================

@router.get("/export/serp")
def export_serp_table(
    project_id: Optional[str] = None,
    engine_ids: Optional[str] = None,
    days: int = Query(30, ge=1, le=365),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    subproject_id: Optional[str] = None,
    prompt_id: Optional[str] = None,
    run_id: Optional[str] = None,
    im_seo_min: Optional[float] = None,
    im_seo_max: Optional[float] = None,
    im_seoia_min: Optional[float] = None,
    im_seoia_max: Optional[float] = None,
    format: str = Query("csv", regex="^(csv|json|excel)$"),
    db: Session = Depends(get_db),
):
    """Exporta Tabela SERP - AI Overview com análise de posicionamento do BB."""
    
    runs, period = _collect_runs(
        db,
        project_id=project_id,
        engine_ids=engine_ids,
        days=days,
        start_date=start_date,
        end_date=end_date,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )
    
    refs = _hydrate_references(db, runs)
    
    # Construir linhas da tabela SERP
    rows = []
    for run in runs:
        prompt = None
        if run.prompt_version_id:
            pv = refs["prompt_versions"].get(run.prompt_version_id)
            if pv:
                prompt = refs["prompts"].get(pv.prompt_id)
        
        subproject = refs["subprojects"].get(run.subproject_id) if run.subproject_id else None
        
        # Parse AI Overview
        serp_feature = run.serp_features[0] if run.serp_features else None
        ai_overview = _parse_ai_overview(serp_feature)
        
        # Construir resposta completa
        resposta = _build_response_text(ai_overview)
        
        # Análise de URLs do BB
        url_info = _get_bb_url_info(ai_overview["references"], run.citations or [])
        
        # Análise de menções ao BB no texto
        prompt_text = prompt.text if prompt else ""
        pergunta_bb = "Sim" if _check_bb_mention(prompt_text) else "Não"
        nome_bb = "Sim" if _check_bb_mention(resposta) else "Não"
        posicao_bb = _find_bb_position(resposta)
        
        # PAA e KP
        tem_paa = "Sim" if serp_feature and serp_feature.has_paa else "Não"
        tem_kp = "Sim" if serp_feature and serp_feature.has_knowledge_panel else "Não"
        
        # Ranking de texto (competitors mencionados)
        ranking_txt = []
        insight_payload = load_insight_payload(run.semantic_insights)
        if insight_payload.get("competitors"):
            for comp in insight_payload["competitors"]:
                ranking_txt.append({
                    "name": comp.get("name"),
                    "mentions": comp.get("mentions", 0)
                })
        
        row = {
            "run_id": run.id,
            "prompt": prompt_text[:500] if prompt_text else "",
            "produto": subproject.name if subproject else "",
            "relevancia": _classify_relevance(run.im_seo_score, run.funnel_stage),
            "pergunta_bb": pergunta_bb,
            "funil": _map_funnel_stage(run.funnel_stage),
            "resposta": resposta[:2000] if resposta else "",
            "tem_url": url_info["tem_url"],
            "url_bb": url_info["url_bb"],
            "posicao_url_bb": url_info["posicao_url_bb"],
            "citation_bb": url_info["citation_bb"],
            "ranking_url": url_info["ranking_url"],
            "nome_bb": nome_bb,
            "posicao_bb": posicao_bb,
            "ranking_txt": json.dumps(ranking_txt, ensure_ascii=False) if ranking_txt else "[]",
            "tem_paa": tem_paa,
            "bb_paa": "Não",  # TODO: Parse paa_questions quando disponível
            "posicao_bb_paa": None,
            "tem_kp": tem_kp,
            "bb_kp": "Não",  # TODO: Parse knowledge_panel_json quando disponível
        }
        rows.append(row)
    
    # Retornar no formato solicitado
    if format == "json":
        return {
            "metadata": _build_metadata(runs=runs, period=period, filters=_build_filters_dict(
                project_id=project_id, engine_ids=engine_ids, days=days
            )),
            "data": rows
        }
    elif format == "csv":
        output = io.StringIO()
        if rows:
            writer = csv.DictWriter(output, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        
        filename = f"serp_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    else:  # excel
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "SERP - AI Overview"
            _populate_sheet(ws, rows)
            
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            filename = f"serp_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return Response(
                content=buffer.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}"}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="Excel export requires 'openpyxl'")


# ============================================================================
# EXPORT ENDPOINT: TABELA INDICADORES - AI OVERVIEW
# ============================================================================

@router.get("/export/indicadores")
def export_indicadores_table(
    project_id: Optional[str] = None,
    engine_ids: Optional[str] = None,
    days: int = Query(30, ge=1, le=365),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    subproject_id: Optional[str] = None,
    prompt_id: Optional[str] = None,
    run_id: Optional[str] = None,
    im_seo_min: Optional[float] = None,
    im_seo_max: Optional[float] = None,
    im_seoia_min: Optional[float] = None,
    im_seoia_max: Optional[float] = None,
    format: str = Query("csv", regex="^(csv|json|excel)$"),
    db: Session = Depends(get_db),
):
    """Exporta Tabela Indicadores - AI Overview com métricas EEAT e percepção."""
    
    runs, period = _collect_runs(
        db,
        project_id=project_id,
        engine_ids=engine_ids,
        days=days,
        start_date=start_date,
        end_date=end_date,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )
    
    refs = _hydrate_references(db, runs)
    
    rows = []
    for run in runs:
        prompt = None
        if run.prompt_version_id:
            pv = refs["prompt_versions"].get(run.prompt_version_id)
            if pv:
                prompt = refs["prompts"].get(pv.prompt_id)
        
        subproject = refs["subprojects"].get(run.subproject_id) if run.subproject_id else None
        prompt_text = prompt.text if prompt else ""
        
        # Verificar menção ao BB
        serp_feature = run.serp_features[0] if run.serp_features else None
        ai_overview = _parse_ai_overview(serp_feature)
        resposta = _build_response_text(ai_overview)
        nome_bb = "Sim" if _check_bb_mention(resposta) else "Não"
        
        # Scores EEAT
        score_eeat = run.eeat_score if run.eeat_score is not None else 0
        experiencia = run.eeat_experience if run.eeat_experience is not None else 0
        expertise = run.eeat_expertise if run.eeat_expertise is not None else 0
        autoridade = run.eeat_authoritativeness if run.eeat_authoritativeness is not None else 0
        confiabilidade = run.eeat_trustworthiness if run.eeat_trustworthiness is not None else 0
        
        # Percepção
        perception_scores = {"inovador": 0, "seguranca": 0, "custo": 0, "atendimento": 0}
        insight_payload = load_insight_payload(run.semantic_insights)
        perception_data = insight_payload.get("perception")
        if perception_data:
            perception_scores = _calculate_perception_scores(perception_data)
        
        # Entidades
        qt_total_entidades = run.entities_detected or 0
        qt_entidades_bb = 0
        entities_payload = insight_payload.get("entities") if 'insight_payload' in locals() else []
        if entities_payload:
            for entity in entities_payload:
                if _check_bb_mention(entity.get("name", "")):
                    qt_entidades_bb += 1
        
        row = {
            "run_id": run.id,
            "prompt": prompt_text[:500] if prompt_text else "",
            "produto": subproject.name if subproject else "",
            "nome_bb": nome_bb,
            "score_eeat_bb": round(score_eeat, 2),
            "experiencia_bb": round(experiencia, 2),
            "expertise_bb": round(expertise, 2),
            "autoridade_bb": round(autoridade, 2),
            "confiabilidade_bb": round(confiabilidade, 2),
            "ranking_eeat": "[]",  # TODO: Implementar ranking cross-run
            "inovador_bb": perception_scores["inovador"],
            "seguranca_bb": perception_scores["seguranca"],
            "custo_bb": perception_scores["custo"],
            "atendimento_bb": perception_scores["atendimento"],
            "resultado_percepcao": json.dumps(perception_scores, ensure_ascii=False),
            "qt_total_entidades": qt_total_entidades,
            "qt_entidades_bb": qt_entidades_bb,
            "conexoes_bancos": "[]",  # TODO: Parse relationships quando disponível
        }
        rows.append(row)
    
    # Retornar no formato solicitado
    if format == "json":
        return {
            "metadata": _build_metadata(runs=runs, period=period, filters=_build_filters_dict(
                project_id=project_id, engine_ids=engine_ids, days=days
            )),
            "data": rows
        }
    elif format == "csv":
        output = io.StringIO()
        if rows:
            writer = csv.DictWriter(output, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        
        filename = f"indicadores_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    else:  # excel
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Indicadores - AI Overview"
            _populate_sheet(ws, rows)
            
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            filename = f"indicadores_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return Response(
                content=buffer.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}"}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="Excel export requires 'openpyxl'")


# ============================================================================
# EXPORT ENDPOINT: TABELA ESTRUTURA WEB - AI OVERVIEW
# ============================================================================

@router.get("/export/estrutura-web")
def export_estrutura_web_table(
    project_id: Optional[str] = None,
    engine_ids: Optional[str] = None,
    days: int = Query(30, ge=1, le=365),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    subproject_id: Optional[str] = None,
    prompt_id: Optional[str] = None,
    run_id: Optional[str] = None,
    im_seo_min: Optional[float] = None,
    im_seo_max: Optional[float] = None,
    im_seoia_min: Optional[float] = None,
    im_seoia_max: Optional[float] = None,
    format: str = Query("csv", regex="^(csv|json|excel)$"),
    db: Session = Depends(get_db),
):
    """Exporta Tabela Estrutura Web - AI Overview com dados de estrutura e performance."""
    
    runs, period = _collect_runs(
        db,
        project_id=project_id,
        engine_ids=engine_ids,
        days=days,
        start_date=start_date,
        end_date=end_date,
        subproject_id=subproject_id,
        prompt_id=prompt_id,
        run_id=run_id,
        im_seo_min=im_seo_min,
        im_seo_max=im_seo_max,
        im_seoia_min=im_seoia_min,
        im_seoia_max=im_seoia_max,
    )
    
    refs = _hydrate_references(db, runs)
    
    rows = []
    for run in runs:
        prompt = None
        if run.prompt_version_id:
            pv = refs["prompt_versions"].get(run.prompt_version_id)
            if pv:
                prompt = refs["prompts"].get(pv.prompt_id)
        
        subproject = refs["subprojects"].get(run.subproject_id) if run.subproject_id else None
        prompt_text = prompt.text if prompt else ""
        
        # URL BB citada?
        has_bb_url = any(cite.is_ours for cite in (run.citations or []))
        url_bb = "Sim" if has_bb_url else "Não"
        
        # AI Ready Blocks
        tem_aiblocks = "Sim" if (run.ia_ready_blocks_count or 0) > 0 else "Não"
        qt_aiblocks = run.ia_ready_blocks_count or 0
        tipos = []
        if run.has_lists:
            tipos.append("listas")
        if run.has_faqs:
            tipos.append("FAQs")
        if run.has_tables:
            tipos.append("tabelas")
        if run.has_step_by_step:
            tipos.append("passo a passo")
        tipos_aiblocks = ", ".join(tipos) if tipos else ""
        
        # Performance
        performance = run.core_web_vitals_score if run.core_web_vitals_score is not None else None
        latencia = run.latency_ms if run.latency_ms is not None else None
        
        row = {
            "run_id": run.id,
            "prompt": prompt_text[:500] if prompt_text else "",
            "produto": subproject.name if subproject else "",
            "url_bb": url_bb,
            "tem_titulo_bb": None,  # Não disponível - requer scraping
            "txt_titulo_bb": None,
            "tem_descricao_bb": None,
            "txt_descricao_bb": None,
            "tem_keywords": None,
            "txt_keywords": None,
            "tem_robots_bb": None,
            "txt_robots_bb": None,
            "tem_ogtags_bb": None,
            "ogtags_bb": None,
            "tem_aiblocks_bb": tem_aiblocks,
            "qt_aiblocks_bb": qt_aiblocks,
            "tipos_aiblocks_bb": tipos_aiblocks,
            "ranking_web": "[]",  # TODO: Implementar ranking cross-run
            "latencia": latencia,
            "performance": round(performance, 2) if performance is not None else None,
        }
        rows.append(row)
    
    # Retornar no formato solicitado
    if format == "json":
        return {
            "metadata": _build_metadata(runs=runs, period=period, filters=_build_filters_dict(
                project_id=project_id, engine_ids=engine_ids, days=days
            )),
            "data": rows
        }
    elif format == "csv":
        output = io.StringIO()
        if rows:
            writer = csv.DictWriter(output, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        
        filename = f"estrutura_web_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
        return Response(
            content=output.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    else:  # excel
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Estrutura Web - AI Overview"
            _populate_sheet(ws, rows)
            
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            filename = f"estrutura_web_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return Response(
                content=buffer.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}"}
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="Excel export requires 'openpyxl'")
